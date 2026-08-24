from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from io import BytesIO
from math import floor
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from db import normalize_week_start_iso
from rules import suggest_taxa_erros_pct
from utils import monday_of, strip_embedded_justification_block


WEEKLY_EVAL_SHEET_NAME = "Avaliacoes"
WEEKLY_EVAL_INSTRUCTIONS_SHEET_NAME = "Instrucoes"
WEEKLY_EVAL_LISTS_SHEET_NAME = "Listas"

WEEKLY_JUSTIFICATION_MODEL_OPTIONS = [
    "Resultado atual",
    "Padrão 100%",
    "Revisão pontual",
    "Acompanhamento",
    "Crítico",
]

WEEKLY_JUSTIFICATION_MODELS = {
    "assiduidade": {
        "excelente": "Manteve assiduidade plena no período, sem faltas, atrasos ou saídas antecipadas que impactassem a rotina operacional.",
        "adequado": "Manteve assiduidade adequada, com eventual ajuste pontual controlado e sem impacto relevante para a operação.",
        "atencao": "Apresentou ocorrência pontual de assiduidade no período, exigindo acompanhamento para evitar reincidência.",
        "critico": "Teve desvios relevantes de assiduidade, com impacto na rotina e necessidade de alinhamento imediato.",
    },
    "qualidade": {
        "excelente": "Executou as atividades com padrão consistente de qualidade, sem retrabalho ou divergência relevante registrada.",
        "adequado": "Manteve qualidade adequada na execução, com pequenos ajustes pontuais dentro do esperado para a operação.",
        "atencao": "Apresentou desvios de qualidade que exigiram correção e reforço de atenção aos procedimentos.",
        "critico": "A qualidade ficou abaixo do esperado, com necessidade de acompanhamento próximo e plano de correção.",
    },
    "taxa_erros": {
        "excelente": "Não houve registro relevante no log de erros, mantendo desempenho compatível com o padrão esperado.",
        "adequado": "Houve ocorrência pontual controlada, sem impacto significativo no resultado geral da semana.",
        "atencao": "A taxa de erros foi impactada por ocorrências no período, exigindo reforço em conferência e prevenção de reincidência.",
        "critico": "Os erros registrados impactaram significativamente o resultado, exigindo ação corretiva imediata e acompanhamento.",
    },
    "produtividade": {
        "excelente": "Manteve ritmo produtivo consistente, com boa aderência ao fluxo e ao volume esperado para a semana.",
        "adequado": "Apresentou produtividade adequada, com oscilação pontual sem prejuízo relevante ao fluxo operacional.",
        "atencao": "A produtividade apresentou queda ou instabilidade no período, exigindo acompanhamento dos gargalos e rotina de execução.",
        "critico": "A produtividade ficou abaixo do esperado, com impacto no fluxo e necessidade de plano de recuperação.",
    },
    "comportamento": {
        "excelente": "Manteve postura adequada, colaborativa e aderente aos procedimentos e à disciplina operacional.",
        "adequado": "Apresentou comportamento geral adequado, com pontos pontuais de ajuste sem impacto relevante na equipe.",
        "atencao": "Foram observados pontos de comportamento que exigem alinhamento, reforço de disciplina e melhoria de comunicação.",
        "critico": "O comportamento no período exigiu intervenção, com necessidade de alinhamento imediato e acompanhamento próximo.",
    },
}

WEEKLY_EVAL_EXPORT_HEADERS = [
    "Importar?",
    "employee_id",
    "Nome",
    "Setor",
    "Função",
    "Semana",
    "Itens",
    "Assiduidade (%)",
    "Qualidade (%)",
    "Taxa Erros (%)",
    "Prod/Efic (%)",
    "Comportamento (%)",
    "Score (%)",
    "Avaliador",
    "Notas",
]

WEEKLY_EVAL_COLUMN_ALIASES = {
    "importar": "importar",
    "importar sim nao": "importar",
    "employee id": "employee_id",
    "id": "employee_id",
    "nome": "nome",
    "funcionario": "nome",
    "colaborador": "nome",
    "setor": "setor",
    "funcao": "funcao",
    "semana": "week_start",
    "week start": "week_start",
    "data da semana": "week_start",
    "itens": "items_count",
    "pecas": "items_count",
    "itens pecas": "items_count",
    "assiduidade": "assiduidade_pct",
    "assiduidade pct": "assiduidade_pct",
    "assiduidade percentual": "assiduidade_pct",
    "qualidade": "qualidade_pct",
    "qualidade pct": "qualidade_pct",
    "qualidade percentual": "qualidade_pct",
    "taxa erros": "taxa_erros_pct",
    "taxa de erros": "taxa_erros_pct",
    "taxa erros pct": "taxa_erros_pct",
    "taxa de erros pct": "taxa_erros_pct",
    "prod efic": "produtividade_pct",
    "prod efic pct": "produtividade_pct",
    "produtividade": "produtividade_pct",
    "produtividade eficiencia": "produtividade_pct",
    "comportamento": "comportamento_pct",
    "comportamento pct": "comportamento_pct",
    "avaliador": "evaluator",
    "notas": "notes",
    "observacao": "notes",
    "observacoes": "notes",
    "score": "score",
    "score pct": "score",
}

PCT_FIELDS = {
    "assiduidade_pct": "Assiduidade (%)",
    "qualidade_pct": "Qualidade (%)",
    "produtividade_pct": "Prod/Efic (%)",
    "comportamento_pct": "Comportamento (%)",
}

ERROR_RATE_STEP = 5


def _safe_float(value: Any, default: float = 100.0) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return float(default)


def round_percentage_to_step(value: float, step: int = ERROR_RATE_STEP) -> float:
    if int(step) <= 0:
        raise ValueError("O intervalo de arredondamento deve ser positivo.")
    step = int(step)
    bounded = min(100.0, max(0.0, float(value)))
    rounded = floor((bounded + (step / 2.0)) / step) * step
    return float(min(100, max(0, rounded)))


def calculate_import_error_rate(
    role: str,
    items_count: int,
    weekly_errors_rows: list[dict] | None,
    step: int = ERROR_RATE_STEP,
) -> dict:
    error_rows = list(weekly_errors_rows or [])
    suggestion = suggest_taxa_erros_pct(
        role=str(role or ""),
        items_count=int(items_count or 0),
        weekly_errors_rows=error_rows,
        strict_critical_zero=True,
        factor=12.0,
    )
    raw_pct = float(suggestion.suggested_pct)
    rounded_pct = round_percentage_to_step(raw_pct, step)
    error_qty = sum(max(0, int(row.get("qty") or 1)) for row in error_rows)
    reason = str(suggestion.reason or "").strip()
    if rounded_pct != raw_pct:
        reason += f" Resultado bruto: {raw_pct:.1f}%; ajustado para {rounded_pct:.0f}% em intervalos de {step}%."
    else:
        reason += f" Resultado mantido em {rounded_pct:.0f}%, já compatível com intervalos de {step}%."
    return {
        "pct": rounded_pct,
        "raw_pct": raw_pct,
        "error_qty": error_qty,
        "reason": reason.strip(),
    }


def template_tier_from_pct(pct: float) -> str:
    value = _safe_float(pct, 100)
    if value >= 91:
        return "excelente"
    if value >= 81:
        return "adequado"
    if value >= 71:
        return "atencao"
    return "critico"


def build_criterion_template_justifications(
    pcts: dict | None = None,
    model: str = "Resultado atual",
) -> dict:
    if model not in WEEKLY_JUSTIFICATION_MODEL_OPTIONS:
        raise ValueError("Modelo de justificativa inválido.")

    tier_by_model = {
        "Padrão 100%": "excelente",
        "Revisão pontual": "adequado",
        "Acompanhamento": "atencao",
        "Crítico": "critico",
    }

    pcts = pcts or {}
    out = {}
    for key, templates in WEEKLY_JUSTIFICATION_MODELS.items():
        tier = tier_by_model.get(model)
        if tier is None:
            tier = template_tier_from_pct(pcts.get(key, 100))
        out[key] = templates[tier]
    return out


def _normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^a-zA-Z0-9]+", " ", text).strip().lower()
    return text


def _canonicalize_import_columns(raw_df: pd.DataFrame) -> pd.DataFrame:
    rename = {}
    used = set()
    for column in raw_df.columns:
        normalized = _normalize_header(column)
        canonical = WEEKLY_EVAL_COLUMN_ALIASES.get(normalized)
        if canonical and canonical not in used:
            rename[column] = canonical
            used.add(canonical)
    return raw_df.rename(columns=rename).copy()


def _excel_text(value: Any) -> str:
    text = _clean_text(value).replace("\r", " ").replace("\n", " ")
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _clean_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    return str(value).strip()


def _parse_percentage(value: Any, label: str) -> tuple[float | None, str]:
    cleaned = _clean_text(value)
    if not cleaned:
        return None, f"{label} vazio"

    cleaned = cleaned.replace("%", "").replace(" ", "")
    if "," in cleaned and "." not in cleaned:
        cleaned = cleaned.replace(",", ".")
    try:
        parsed = float(cleaned)
    except (TypeError, ValueError):
        return None, f"{label} inválido"

    if parsed < 0 or parsed > 100:
        return None, f"{label} fora de 0–100"
    return round(parsed, 2), ""


def _parse_non_negative_int(value: Any, label: str) -> tuple[int | None, str]:
    cleaned = _clean_text(value)
    if not cleaned:
        return 0, ""
    try:
        numeric = float(cleaned.replace(",", "."))
    except (TypeError, ValueError):
        return None, f"{label} inválido"
    if numeric < 0 or not numeric.is_integer():
        return None, f"{label} deve ser um inteiro não negativo"
    return int(numeric), ""


def _parse_employee_id(value: Any) -> tuple[int | None, str]:
    parsed, problem = _parse_non_negative_int(value, "employee_id")
    if problem:
        return None, problem
    if not parsed:
        return None, "employee_id vazio"
    return parsed, ""


def _parse_import_flag(value: Any) -> tuple[bool | None, str]:
    normalized = _normalize_header(_clean_text(value))
    if normalized in {"sim", "s", "yes", "y", "1", "true", "verdadeiro"}:
        return True, ""
    if normalized in {"nao", "n", "no", "0", "false", "falso", ""}:
        return False, ""
    return None, "Importar? deve ser SIM ou NÃO"


def _parse_week(value: Any) -> tuple[str | None, str]:
    if isinstance(value, pd.Timestamp):
        value = value.to_pydatetime()
    try:
        parsed = datetime.strptime(normalize_week_start_iso(value), "%Y-%m-%d").date()
        return monday_of(parsed).isoformat(), ""
    except Exception:
        return None, "Semana inválida"


def weekly_eval_import_week_starts(raw_df: pd.DataFrame) -> list[str]:
    data = _canonicalize_import_columns(raw_df)
    if "week_start" not in data.columns:
        return []

    weeks = set()
    for _, row in data.iterrows():
        should_import, flag_problem = _parse_import_flag(row.get("importar"))
        if flag_problem or should_import is not True:
            continue
        week_start_iso, problem = _parse_week(row.get("week_start"))
        if not problem and week_start_iso:
            weeks.add(week_start_iso)
    return sorted(weeks)


def _build_notes_with_justifications(notes: str, justifications: dict) -> str:
    notes_clean = strip_embedded_justification_block(
        str(notes or "").strip(),
        "JUSTIFICATIVAS (SEMANA)",
    )
    block = (
        "JUSTIFICATIVAS (SEMANA)\n"
        f"- Assiduidade: {justifications['assiduidade']}\n"
        f"- Qualidade: {justifications['qualidade']}\n"
        f"- Taxa de Erros: {justifications['taxa_erros']}\n"
        f"- Produtividade/Eficiência: {justifications['produtividade']}\n"
        f"- Comportamento: {justifications['comportamento']}\n"
    )
    return (notes_clean + "\n\n" + block).strip()


def build_weekly_eval_workbook_bytes(
    mass_df: pd.DataFrame,
    week_start_iso: str,
    evaluator_options: list[str] | None = None,
) -> bytes:
    week_start_iso = normalize_week_start_iso(week_start_iso)
    week_start = datetime.strptime(week_start_iso, "%Y-%m-%d").date()
    evaluator_options = [str(value).strip() for value in (evaluator_options or []) if str(value).strip()]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = WEEKLY_EVAL_SHEET_NAME
    instructions = workbook.create_sheet(WEEKLY_EVAL_INSTRUCTIONS_SHEET_NAME)
    lists = workbook.create_sheet(WEEKLY_EVAL_LISTS_SHEET_NAME)

    navy = "17324D"
    blue = "1F6F8B"
    pale_blue = "DCEAF2"
    pale_yellow = "FFF2CC"
    pale_gray = "E7E6E6"
    pale_red = "F4CCCC"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="C8D0D8")

    for col_idx, header in enumerate(WEEKLY_EVAL_EXPORT_HEADERS, start=1):
        cell = sheet.cell(row=1, column=col_idx, value=header)
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=blue))

    selected_mask = mass_df.get("Selecionar", pd.Series(False, index=mass_df.index)).fillna(False).astype(bool)
    has_explicit_selection = bool(selected_mask.any())

    for excel_row, (_, row) in enumerate(mass_df.iterrows(), start=2):
        should_import = bool(row.get("Selecionar", False)) if has_explicit_selection else True
        values = [
            "SIM" if should_import else "NÃO",
            int(row["employee_id"]),
            _excel_text(row.get("Nome", "")),
            _excel_text(row.get("Setor", "")),
            _excel_text(row.get("Função", "")),
            week_start,
            int(_safe_float(row.get("Itens", 0), 0)),
            _safe_float(row.get("Assiduidade (%)", 100), 100),
            _safe_float(row.get("Qualidade (%)", 100), 100),
            None,
            _safe_float(row.get("Prod/Efic (%)", 100), 100),
            _safe_float(row.get("Comportamento (%)", 100), 100),
            f'=IF(COUNTA(H{excel_row}:L{excel_row})<5,"",ROUND(AVERAGE(H{excel_row}:L{excel_row}),1))',
            _excel_text(row.get("Avaliador", "")),
            str(row.get("Notas", "") or ""),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=col_idx, value=value)
            cell.border = Border(bottom=thin_gray)
            cell.alignment = Alignment(vertical="top", wrap_text=col_idx == 15)

        for col_idx in (2, 3, 4, 5, 13):
            sheet.cell(excel_row, col_idx).fill = PatternFill("solid", fgColor=pale_gray)
            sheet.cell(excel_row, col_idx).protection = Protection(locked=True)
        sheet.cell(excel_row, 10).fill = PatternFill("solid", fgColor=pale_blue)
        sheet.cell(excel_row, 10).protection = Protection(locked=True)
        for col_idx in (1, 6, 7, 8, 9, 11, 12, 14, 15):
            sheet.cell(excel_row, col_idx).fill = PatternFill("solid", fgColor=pale_yellow)
            sheet.cell(excel_row, col_idx).protection = Protection(locked=False)

        sheet.cell(excel_row, 6).number_format = "yyyy-mm-dd"
        sheet.cell(excel_row, 7).number_format = "#,##0"
        for col_idx in range(8, 14):
            sheet.cell(excel_row, col_idx).number_format = "0.0"

    max_row = max(2, sheet.max_row)
    table_ref = f"A1:O{max_row}"
    table = Table(displayName="TabelaAvaliacoes", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = table_ref
    sheet.row_dimensions[1].height = 34
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 13
    sheet.column_dimensions["B"].width = 13
    sheet.column_dimensions["C"].width = 30
    sheet.column_dimensions["D"].width = 20
    sheet.column_dimensions["E"].width = 22
    sheet.column_dimensions["F"].width = 14
    sheet.column_dimensions["G"].width = 12
    for column in ("H", "I", "J", "K", "L", "M"):
        sheet.column_dimensions[column].width = 18
    sheet.column_dimensions["N"].width = 26
    sheet.column_dimensions["O"].width = 42

    sheet["B1"].comment = Comment("Identificador interno. Não altere.", "Sistema")
    sheet["F1"].comment = Comment(
        "Informe qualquer data da semana desejada. O app normaliza a data para a segunda-feira correspondente.",
        "Sistema",
    )
    sheet["J1"].comment = Comment(
        "Deixe esta coluna vazia. O app calcula a taxa com os erros já registrados para o funcionário e a semana.",
        "Sistema",
    )
    sheet["M1"].comment = Comment(
        "O score final será exibido na prévia depois que o app calcular a Taxa de Erros.",
        "Sistema",
    )

    import_validation = DataValidation(type="list", formula1='"SIM,NÃO"', allow_blank=False)
    import_validation.error = "Escolha SIM ou NÃO."
    import_validation.errorTitle = "Valor inválido"
    import_validation.prompt = "Somente as linhas marcadas como SIM serão importadas."
    import_validation.promptTitle = "Selecionar linha"
    import_validation.showErrorMessage = True
    import_validation.showInputMessage = True
    sheet.add_data_validation(import_validation)
    import_validation.add(f"A2:A{max_row}")

    pct_validation = DataValidation(
        type="decimal",
        operator="between",
        formula1="0",
        formula2="100",
        allow_blank=False,
    )
    pct_validation.error = "Informe um número entre 0 e 100."
    pct_validation.errorTitle = "Percentual inválido"
    pct_validation.showErrorMessage = True
    sheet.add_data_validation(pct_validation)
    pct_validation.add(f"H2:I{max_row}")
    pct_validation.add(f"K2:L{max_row}")

    items_validation = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
    items_validation.error = "Informe um número inteiro maior ou igual a zero."
    items_validation.errorTitle = "Quantidade inválida"
    items_validation.showErrorMessage = True
    sheet.add_data_validation(items_validation)
    items_validation.add(f"G2:G{max_row}")

    week_validation = DataValidation(
        type="date",
        operator="between",
        formula1="DATE(2020,1,1)",
        formula2="DATE(2100,12,31)",
        allow_blank=False,
    )
    week_validation.error = "Informe uma data válida para a semana."
    week_validation.errorTitle = "Data inválida"
    week_validation.showErrorMessage = True
    sheet.add_data_validation(week_validation)
    week_validation.add(f"F2:F{max_row}")

    lists.append(["Avaliadores ativos"])
    for evaluator in evaluator_options:
        lists.append([_excel_text(evaluator)])
    if evaluator_options:
        workbook.defined_names.add(DefinedName(
            "Avaliadores",
            attr_text=f"'{WEEKLY_EVAL_LISTS_SHEET_NAME}'!$A$2:$A${len(evaluator_options) + 1}",
        ))
        evaluator_validation = DataValidation(
            type="list",
            formula1="=Avaliadores",
            allow_blank=True,
        )
        evaluator_validation.error = "Selecione um avaliador da lista."
        evaluator_validation.errorTitle = "Avaliador inválido"
        evaluator_validation.showErrorMessage = True
        sheet.add_data_validation(evaluator_validation)
        evaluator_validation.add(f"N2:N{max_row}")
    lists.sheet_state = "hidden"

    sheet.conditional_formatting.add(
        f"A2:A{max_row}",
        FormulaRule(formula=["$A2=\"SIM\""], fill=PatternFill("solid", fgColor="D9EAD3")),
    )
    sheet.conditional_formatting.add(
        f"A2:A{max_row}",
        FormulaRule(formula=["$A2<>\"SIM\""], fill=PatternFill("solid", fgColor=pale_red)),
    )

    instructions.sheet_view.showGridLines = False
    instructions["A1"] = "Avaliação semanal em massa — guia rápido"
    instructions["A1"].fill = PatternFill("solid", fgColor=navy)
    instructions["A1"].font = Font(color=white, bold=True, size=16)
    instructions["A1"].alignment = Alignment(vertical="center")
    instructions.merge_cells("A1:F2")
    instructions.row_dimensions[1].height = 24
    instruction_rows = [
        (4, "1. Escolha as linhas", "Use SIM em “Importar?” somente para os colaboradores que devem ser gravados."),
        (6, "2. Preencha semana e resultados", "A coluna Semana aceita qualquer data e é normalizada para a segunda-feira correspondente. Informe Itens, Assiduidade, Qualidade, Prod/Efic e Comportamento. Deixe Taxa Erros (%) vazia."),
        (8, "3. Defina o avaliador", "Escolha um avaliador ativo na lista. Se deixar vazio, o app usará o avaliador padrão escolhido na importação."),
        (10, "4. Importe no app", "Envie este XLSX na Avaliação em massa, escolha o modelo de justificativa e gere a prévia."),
        (12, "5. Confirme a gravação", "Para cada linha, o app consulta os erros da semana informada, calcula a taxa em intervalos de 5% e mostra as justificativas antes de atualizar o banco."),
    ]
    for row_idx, title, detail in instruction_rows:
        instructions.cell(row_idx, 1, title)
        instructions.cell(row_idx, 1).font = Font(bold=True, color=blue, size=12)
        instructions.cell(row_idx + 1, 1, detail)
        instructions.cell(row_idx + 1, 1).alignment = Alignment(wrap_text=True, vertical="top")
        instructions.row_dimensions[row_idx + 1].height = 32
        instructions.merge_cells(start_row=row_idx, start_column=1, end_row=row_idx, end_column=6)
        instructions.merge_cells(start_row=row_idx + 1, start_column=1, end_row=row_idx + 1, end_column=6)
    instructions["A15"] = "Importante"
    instructions["A15"].font = Font(bold=True, color="9C0006")
    instructions["A16"] = (
        "A importação atualiza a avaliação que já existir para o mesmo employee_id e semana. "
        "A Taxa de Erros é sempre recalculada com o log semanal; qualquer valor digitado nessa coluna será ignorado. "
        "As cinco justificativas são geradas no app pelo modelo escolhido."
    )
    instructions["A16"].fill = PatternFill("solid", fgColor="FCE8E6")
    instructions["A16"].alignment = Alignment(wrap_text=True, vertical="top")
    instructions.merge_cells("A15:F15")
    instructions.merge_cells("A16:F18")
    instructions.column_dimensions["A"].width = 24
    for column in ("B", "C", "D", "E", "F"):
        instructions.column_dimensions[column].width = 16
    instructions.row_dimensions[16].height = 45

    workbook.active = workbook.sheetnames.index(WEEKLY_EVAL_SHEET_NAME)
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except Exception:
        pass

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def load_weekly_eval_import_file(uploaded_file) -> pd.DataFrame:
    if hasattr(uploaded_file, "seek"):
        uploaded_file.seek(0)
    try:
        return pd.read_excel(uploaded_file, sheet_name=WEEKLY_EVAL_SHEET_NAME, dtype=object)
    except ValueError as exc:
        raise ValueError(f"Não encontrei a aba {WEEKLY_EVAL_SHEET_NAME} no XLSX.") from exc


def prepare_weekly_eval_import(
    raw_df: pd.DataFrame,
    employees_df: pd.DataFrame,
    evaluator_options: list[str],
    default_week_start_iso: str,
    justification_model: str,
    default_evaluator: str = "",
    weekly_errors_by_key: dict[tuple[int, str], list[dict]] | None = None,
) -> dict:
    if justification_model not in WEEKLY_JUSTIFICATION_MODEL_OPTIONS:
        raise ValueError("Selecione um modelo de justificativa válido.")

    data = _canonicalize_import_columns(raw_df)
    required = {"importar", "employee_id", "week_start", *PCT_FIELDS.keys()}
    missing = sorted(required - set(data.columns))
    if missing:
        raise ValueError("Colunas obrigatórias ausentes: " + ", ".join(missing))

    default_week_start_iso = normalize_week_start_iso(default_week_start_iso)
    evaluator_options = [str(value).strip() for value in evaluator_options if str(value).strip()]
    evaluator_lookup = {value.casefold(): value for value in evaluator_options}
    default_evaluator = str(default_evaluator or "").strip()
    weekly_errors_by_key = weekly_errors_by_key or {}

    employees = employees_df.copy()
    employee_lookup = {
        int(row["id"]): row
        for _, row in employees.iterrows()
        if row.get("id") is not None
    }

    preview_rows = []
    valid_rows = []
    seen_keys = set()
    total = 0
    selected = 0
    ignored = 0
    invalid = 0

    for position, row in data.iterrows():
        excel_row = int(position) + 2 if isinstance(position, int) else len(preview_rows) + 2
        if all(pd.isna(value) or str(value).strip() == "" for value in row.tolist()):
            continue

        total += 1
        should_import, flag_problem = _parse_import_flag(row.get("importar"))
        if should_import is False and not flag_problem:
            ignored += 1
            continue

        selected += 1
        problems = [flag_problem] if flag_problem else []

        employee_id, employee_problem = _parse_employee_id(row.get("employee_id"))
        if employee_problem:
            problems.append(employee_problem)
        employee = employee_lookup.get(employee_id) if employee_id is not None else None
        if employee_id is not None and employee is None:
            problems.append("employee_id não pertence a um colaborador ativo e avaliável")

        week_start_iso, week_problem = _parse_week(row.get("week_start"))
        if week_problem:
            problems.append(week_problem)
        key = (employee_id, week_start_iso)
        if employee_id is not None and week_start_iso and key in seen_keys:
            problems.append("colaborador duplicado no arquivo para a mesma semana")
        seen_keys.add(key)

        parsed_pcts = {}
        for field, label in PCT_FIELDS.items():
            value, problem = _parse_percentage(row.get(field), label)
            if problem:
                problems.append(problem)
            parsed_pcts[field] = value

        items_count, items_problem = _parse_non_negative_int(row.get("items_count", 0), "Itens")
        if items_problem:
            problems.append(items_problem)

        evaluator_raw = _clean_text(row.get("evaluator", "")) or default_evaluator
        evaluator = evaluator_lookup.get(evaluator_raw.casefold(), "") if evaluator_raw else ""
        if not evaluator:
            problems.append("Avaliador vazio ou fora da lista de avaliadores ativos")

        employee_name = _clean_text(employee.get("name", "")) if employee is not None else _clean_text(row.get("nome", ""))
        if problems:
            invalid += 1
            preview_rows.append({
                "Linha": excel_row,
                "Status": "REVISAR",
                "Funcionário": employee_name,
                "Semana": week_start_iso or "",
                "Avaliador": evaluator_raw,
                "Modelo": justification_model,
                "Mensagem": "; ".join(problems),
            })
            continue

        pcts = {
            "assiduidade": parsed_pcts["assiduidade_pct"],
            "qualidade": parsed_pcts["qualidade_pct"],
            "produtividade": parsed_pcts["produtividade_pct"],
            "comportamento": parsed_pcts["comportamento_pct"],
        }
        error_rows = weekly_errors_by_key.get((int(employee_id), str(week_start_iso)), [])
        error_rate = calculate_import_error_rate(
            role=_clean_text(employee.get("role", "")),
            items_count=int(items_count),
            weekly_errors_rows=error_rows,
        )
        pcts["taxa_erros"] = float(error_rate["pct"])
        justifications = build_criterion_template_justifications(pcts, justification_model)
        notes = _clean_text(row.get("notes", ""))

        valid_rows.append({
            "employee_id": int(employee_id),
            "week_start_iso": str(week_start_iso),
            "evaluator": evaluator,
            "notes": _build_notes_with_justifications(notes, justifications),
            "assiduidade_pct": float(pcts["assiduidade"]),
            "qualidade_pct": float(pcts["qualidade"]),
            "taxa_erros_pct": float(pcts["taxa_erros"]),
            "produtividade_pct": float(pcts["produtividade"]),
            "comportamento_pct": float(pcts["comportamento"]),
            "efficiency_pct": float(pcts["produtividade"]),
            "items_count": int(items_count),
            "assiduidade_just": justifications["assiduidade"],
            "qualidade_just": justifications["qualidade"],
            "taxa_erros_just": justifications["taxa_erros"],
            "produtividade_just": justifications["produtividade"],
            "comportamento_just": justifications["comportamento"],
            "taxa_erros_raw_pct": float(error_rate["raw_pct"]),
            "taxa_erros_reason": error_rate["reason"],
            "weekly_error_qty": int(error_rate["error_qty"]),
        })
        preview_rows.append({
            "Linha": excel_row,
            "Status": "OK",
            "Funcionário": employee_name,
            "Semana": week_start_iso,
            "Avaliador": evaluator,
            "Modelo": justification_model,
            "Erros no log": int(error_rate["error_qty"]),
            "Taxa calculada (%)": float(error_rate["pct"]),
            "Regra da taxa": error_rate["reason"],
            "Score (%)": round(sum(pcts.values()) / len(pcts), 1),
            "Assiduidade Just.": justifications["assiduidade"],
            "Qualidade Just.": justifications["qualidade"],
            "Taxa Erros Just.": justifications["taxa_erros"],
            "Produtividade Just.": justifications["produtividade"],
            "Comportamento Just.": justifications["comportamento"],
            "Mensagem": "Pronto para importar",
        })

    return {
        "summary": {
            "total": total,
            "selected": selected,
            "valid": len(valid_rows),
            "invalid": invalid,
            "ignored": ignored,
        },
        "preview_df": pd.DataFrame(preview_rows),
        "valid_rows": valid_rows,
        "week_start_iso": default_week_start_iso,
        "week_start_isos": sorted({row["week_start_iso"] for row in valid_rows}),
        "justification_model": justification_model,
        "default_evaluator": default_evaluator,
    }

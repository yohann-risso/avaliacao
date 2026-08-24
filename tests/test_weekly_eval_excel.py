from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

import db
from weekly_eval_excel import (
    WEEKLY_EVAL_INSTRUCTIONS_SHEET_NAME,
    WEEKLY_EVAL_LISTS_SHEET_NAME,
    WEEKLY_EVAL_SHEET_NAME,
    build_weekly_eval_workbook_bytes,
    calculate_import_error_rate,
    load_weekly_eval_import_file,
    prepare_weekly_eval_import,
)


def _employees() -> pd.DataFrame:
    return pd.DataFrame([
        {
            "id": 10,
            "name": "Ana Expedição",
            "sector": "Expedição",
            "role": "Separador",
        },
        {
            "id": 20,
            "name": "Bruno Estoque",
            "sector": "Estoque",
            "role": "Conferente",
        },
    ])


def _mass_df() -> pd.DataFrame:
    return pd.DataFrame([
        {
            "Selecionar": True,
            "employee_id": 10,
            "Nome": "Ana Expedição",
            "Setor": "Expedição",
            "Função": "Separador",
            "Itens": 120,
            "Assiduidade (%)": 100,
            "Qualidade (%)": 85,
            "Taxa Erros (%)": 75,
            "Prod/Efic (%)": 90,
            "Comportamento (%)": 100,
            "Avaliador": "Líder Um",
            "Notas": "Semana de teste",
        },
        {
            "Selecionar": False,
            "employee_id": 20,
            "Nome": "Bruno Estoque",
            "Setor": "Estoque",
            "Função": "Conferente",
            "Itens": 80,
            "Assiduidade (%)": 100,
            "Qualidade (%)": 100,
            "Taxa Erros (%)": 100,
            "Prod/Efic (%)": 100,
            "Comportamento (%)": 100,
            "Avaliador": "Líder Um",
            "Notas": "",
        },
    ])


def test_build_weekly_eval_workbook_has_editable_template_and_validations():
    content = build_weekly_eval_workbook_bytes(
        _mass_df(),
        "2026-08-17",
        ["Líder Um", "Líder Dois"],
    )

    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames == [
        WEEKLY_EVAL_SHEET_NAME,
        WEEKLY_EVAL_INSTRUCTIONS_SHEET_NAME,
        WEEKLY_EVAL_LISTS_SHEET_NAME,
    ]
    assert workbook[WEEKLY_EVAL_LISTS_SHEET_NAME].sheet_state == "hidden"

    sheet = workbook[WEEKLY_EVAL_SHEET_NAME]
    assert sheet["A2"].value == "SIM"
    assert sheet["A3"].value == "NÃO"
    assert sheet["B2"].value == 10
    assert sheet["F2"].value.date().isoformat() == "2026-08-17"
    assert sheet["F2"].protection.locked is False
    assert sheet["J2"].value is None
    assert sheet["J3"].value is None
    assert sheet["M2"].value == '=IF(COUNTA(H2:L2)<5,"",ROUND(AVERAGE(H2:L2),1))'
    assert "Avaliadores" in workbook.defined_names
    assert {validation.type for validation in sheet.data_validations.dataValidation} == {
        "list",
        "decimal",
        "whole",
        "date",
    }
    decimal_validation = next(
        validation
        for validation in sheet.data_validations.dataValidation
        if validation.type == "decimal"
    )
    assert "J" not in str(decimal_validation.sqref)

    loaded = load_weekly_eval_import_file(BytesIO(content))
    assert list(loaded["employee_id"].astype(int)) == [10, 20]
    assert loaded.iloc[0]["Importar?"] == "SIM"
    preview = prepare_weekly_eval_import(
        loaded,
        _employees(),
        ["Líder Um", "Líder Dois"],
        "2026-08-17",
        "Resultado atual",
        default_evaluator="Líder Um",
    )
    assert preview["summary"] == {
        "total": 2,
        "selected": 1,
        "valid": 1,
        "invalid": 0,
        "ignored": 1,
    }


def test_prepare_weekly_eval_import_generates_justifications_from_selected_model():
    raw = pd.DataFrame([
        {
            "Importar?": "SIM",
            "employee_id": 10,
            "Nome": "Ana Expedição",
            "Semana": "2026-08-17",
            "Itens": 120,
            "Assiduidade (%)": 100,
            "Qualidade (%)": 85,
            "Taxa Erros (%)": pd.NA,
            "Prod/Efic (%)": 90,
            "Comportamento (%)": 70,
            "Avaliador": pd.NA,
            "Notas": "Acompanhamento registrado.",
        },
        {
            "Importar?": "NÃO",
            "employee_id": 20,
            "Semana": "2026-08-17",
            "Itens": 80,
            "Assiduidade (%)": 100,
            "Qualidade (%)": 100,
            "Taxa Erros (%)": 100,
            "Prod/Efic (%)": 100,
            "Comportamento (%)": 100,
            "Avaliador": "Líder Um",
        },
    ])

    preview = prepare_weekly_eval_import(
        raw,
        _employees(),
        ["Líder Um"],
        "2026-08-17",
        "Resultado atual",
        default_evaluator="Líder Um",
        weekly_errors_by_key={
            (10, "2026-08-17"): [
                {"error_type": "Erro de Picking", "severity": "ALTO", "qty": 1},
                {"error_type": "Avaria", "severity": "BAIXO", "qty": 1},
            ],
        },
    )

    assert preview["summary"] == {
        "total": 2,
        "selected": 1,
        "valid": 1,
        "invalid": 0,
        "ignored": 1,
    }
    payload = preview["valid_rows"][0]
    assert payload["employee_id"] == 10
    assert payload["evaluator"] == "Líder Um"
    assert payload["taxa_erros_pct"] == 75
    assert payload["weekly_error_qty"] == 2
    assert "assiduidade plena" in payload["assiduidade_just"]
    assert "qualidade adequada" in payload["qualidade_just"]
    assert "taxa de erros foi impactada" in payload["taxa_erros_just"].lower()
    assert "produtividade adequada" in payload["produtividade_just"].lower()
    assert "intervenção" in payload["comportamento_just"]
    assert payload["notes"].startswith("Acompanhamento registrado.")
    assert "JUSTIFICATIVAS (SEMANA)" in payload["notes"]
    assert preview["preview_df"].iloc[0]["Score (%)"] == 84.0
    assert preview["preview_df"].iloc[0]["Taxa calculada (%)"] == 75
    assert preview["preview_df"].iloc[0]["Erros no log"] == 2


def test_calculate_import_error_rate_rounds_to_nearest_five_percent():
    calculated = calculate_import_error_rate(
        role="Conferente",
        items_count=0,
        weekly_errors_rows=[
            {"error_type": "Avaria", "severity": "MEDIO", "qty": 1},
        ],
    )

    assert calculated["raw_pct"] == 93.3
    assert calculated["pct"] == 95
    assert calculated["pct"] % 5 == 0
    assert "ajustado para 95%" in calculated["reason"]


def test_prepare_weekly_eval_import_accepts_multiple_weeks_from_spreadsheet():
    raw = pd.DataFrame([
        {
            "Importar?": "SIM",
            "employee_id": 10,
            "Semana": "2026-08-19",
            "Itens": 100,
            "Assiduidade (%)": 100,
            "Qualidade (%)": 100,
            "Prod/Efic (%)": 100,
            "Comportamento (%)": 100,
            "Avaliador": "Líder Um",
        },
        {
            "Importar?": "SIM",
            "employee_id": 10,
            "Semana": "2026-08-24",
            "Itens": 100,
            "Assiduidade (%)": 100,
            "Qualidade (%)": 100,
            "Prod/Efic (%)": 100,
            "Comportamento (%)": 100,
            "Avaliador": "Líder Um",
        },
    ])

    preview = prepare_weekly_eval_import(
        raw,
        _employees(),
        ["Líder Um"],
        "2026-08-03",
        "Resultado atual",
        default_evaluator="Líder Um",
        weekly_errors_by_key={
            (10, "2026-08-17"): [
                {"error_type": "Erro de Picking", "severity": "MEDIO", "qty": 1},
            ],
        },
    )

    assert preview["summary"]["valid"] == 2
    assert preview["week_start_isos"] == ["2026-08-17", "2026-08-24"]
    assert [row["week_start_iso"] for row in preview["valid_rows"]] == [
        "2026-08-17",
        "2026-08-24",
    ]
    assert [row["taxa_erros_pct"] for row in preview["valid_rows"]] == [90, 100]


def test_prepare_weekly_eval_import_blocks_wrong_week_duplicates_and_invalid_values():
    raw = pd.DataFrame([
        {
            "Importar?": "SIM",
            "employee_id": 10,
            "Semana": "2026-08-10",
            "Itens": -1,
            "Assiduidade (%)": 101,
            "Qualidade (%)": 80,
            "Taxa Erros (%)": pd.NA,
            "Prod/Efic (%)": 80,
            "Comportamento (%)": 80,
            "Avaliador": "Desconhecido",
        },
        {
            "Importar?": "SIM",
            "employee_id": 10,
            "Semana": "2026-08-10",
            "Itens": 1,
            "Assiduidade (%)": 80,
            "Qualidade (%)": 80,
            "Taxa Erros (%)": pd.NA,
            "Prod/Efic (%)": 80,
            "Comportamento (%)": 80,
            "Avaliador": "Líder Um",
        },
    ])

    preview = prepare_weekly_eval_import(
        raw,
        _employees(),
        ["Líder Um"],
        "2026-08-17",
        "Acompanhamento",
        default_evaluator="Líder Um",
    )

    assert preview["summary"]["valid"] == 0
    assert preview["summary"]["invalid"] == 2
    messages = " | ".join(preview["preview_df"]["Mensagem"].tolist())
    assert "fora de 0–100" in messages
    assert "inteiro não negativo" in messages
    assert "duplicado" in messages
    assert "Avaliador" in messages


def test_import_payload_can_be_persisted_as_weekly_evaluation(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", str(tmp_path / "avaliacoes_excel.db"))
    db.init_db()
    db.insert_employee(
        name="Ana Expedição",
        sector="Expedição",
        role="Separador",
        is_monitor=False,
        hire_date="01/08/2026",
    )
    employee = db.list_active_employees().iloc[0]
    employee_id = int(employee["id"])
    employees = pd.DataFrame([employee])
    raw = pd.DataFrame([{
        "Importar?": "SIM",
        "employee_id": employee_id,
        "Semana": "2026-08-17",
        "Itens": 55,
        "Assiduidade (%)": 100,
        "Qualidade (%)": 80,
        "Taxa Erros (%)": pd.NA,
        "Prod/Efic (%)": 80,
        "Comportamento (%)": 100,
        "Avaliador": "Líder Um",
        "Notas": "Importada do Excel",
    }])

    db.add_weekly_error(
        employee_id=employee_id,
        week_start_iso="2026-08-17",
        role_snapshot="Separador",
        error_type="Erro de Picking",
        severity="ALTO",
        qty=1,
        notes="Erro usado no cálculo automático",
    )
    db.add_weekly_error(
        employee_id=employee_id,
        week_start_iso="2026-08-24",
        role_snapshot="Separador",
        error_type="Avaria",
        severity="BAIXO",
        qty=1,
        notes="Erro de outra semana",
    )
    weekly_errors = db.list_weekly_errors_for_employees_weeks(
        [employee_id],
        ["2026-08-17", "2026-08-24"],
    )
    assert set(weekly_errors["week_start"].tolist()) == {"2026-08-17", "2026-08-24"}
    error_lookup = {
        (int(current_employee_id), str(week_start)): group.to_dict(orient="records")
        for (current_employee_id, week_start), group in weekly_errors.groupby(
            ["employee_id", "week_start"],
            sort=False,
        )
    }

    preview = prepare_weekly_eval_import(
        raw,
        employees,
        ["Líder Um"],
        "2026-08-17",
        "Revisão pontual",
        default_evaluator="Líder Um",
        weekly_errors_by_key=error_lookup,
    )
    db.upsert_weekly_evals(preview["valid_rows"])

    saved = db.get_weekly_eval(employee_id, "2026-08-17").iloc[0]
    assert saved["evaluator"] == "Líder Um"
    assert int(saved["items_count"]) == 55
    assert float(saved["qualidade_pct"]) == 80
    assert float(saved["taxa_erros_pct"]) == 55
    assert "qualidade adequada" in saved["qualidade_just"].lower()
    assert "Importada do Excel" in saved["notes"]
